import os
import pygame
import random
import openpyxl

def charger_record_joueur(nom_joueur):
    BASE_DIR = os.path.dirname(__file__)
    chemin_excel = os.path.join(BASE_DIR, "assets", "records.xlsx")

    if not os.path.exists(chemin_excel):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Joueur"
        ws["B1"] = "Record"
        wb.save(chemin_excel)

    wb = openpyxl.load_workbook(chemin_excel)
    ws = wb.active

    for ligne in ws.iter_rows(min_row=2):
        if ligne[0].value == nom_joueur:
            return ligne[1].value

    return 0

def charger_record_global():
    BASE_DIR = os.path.dirname(__file__)
    chemin_excel = os.path.join(BASE_DIR, "assets", "records.xlsx")

    if not os.path.exists(chemin_excel):
        return 0

    wb = openpyxl.load_workbook(chemin_excel)
    ws = wb.active

    meilleur = 0

    for ligne in ws.iter_rows(min_row=2):
        score = ligne[1].value
        if score is not None and score > meilleur:
            meilleur = score

    return meilleur


def sauvegarder_record_joueur(nom_joueur, nouveau_score):
    BASE_DIR = os.path.dirname(__file__)
    chemin_excel = os.path.join(BASE_DIR, "assets", "records.xlsx")

    if not os.path.exists(chemin_excel):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Joueur"
        ws["B1"] = "Record"
    else:
        wb = openpyxl.load_workbook(chemin_excel)
        ws = wb.active

    joueur_trouve = False

    for ligne in ws.iter_rows(min_row=2):
        if ligne[0].value == nom_joueur:
            joueur_trouve = True
            if nouveau_score > ligne[1].value:
                ligne[1].value = nouveau_score

    if not joueur_trouve:
        ws.append([nom_joueur, nouveau_score])

    wb.save(chemin_excel)


BLANC = (255, 255, 255)
JAUNE = (255, 255, 102)
NOIR = (0, 0, 0)
ROUGE = (213, 50, 80)
VERT = (0, 155, 0)
BLEU = (173, 216, 230)
BLEU2 = (0, 155, 0)
VERT2 = (155, 255, 155)
VERT3 = (0, 100, 0)
Rouge2 = (109, 7, 26)
MARRON = (128, 64, 0)
VERT4 = (0, 255, 0)

VERT_dégradé = (0, 119, 52)
VERT2_dégradé = (0, 155, 67)
VERT3_dégradé = (71, 185, 71)
VERT4_dégradé = (0, 255, 111)

unite = 32
largeur_zone = 60
hauteur_zone = 34

LIM_LEFT = 13
LIM_RIGHT = 46
LIM_TOP = 0
LIM_BOTTOM = hauteur_zone - 1

couleur_nourriture = ROUGE
couleur_serpent = VERT
couleur_fond = VERT2
couleur_coeur = Rouge2


def dessine_carre(x: int, y: int, couleur: tuple, fenetre):
    pygame.draw.rect(
        surface=fenetre,
        color=couleur,
        rect=(x * unite, y * unite, unite, unite)
    )


def dessine_rect(x: int, y: int, couleur: tuple, fenetre):
    pygame.draw.rect(
        surface=fenetre,
        color=couleur,
        rect=(x * unite, y * unite, 417, 1080)
    )

def dessine_bouton_pixel(x, y, largeur, hauteur, couleur, fenetre):
    bord = (
        max(couleur[0] - 70, 0),
        max(couleur[1] - 70, 0),
        max(couleur[2] - 70, 0)
    )

    highlight = (
        min(couleur[0] + 50, 255),
        min(couleur[1] + 50, 255),
        min(couleur[2] + 50, 255)
    )
    

    pixel = 8
    cols = largeur // pixel
    rows = hauteur // pixel

    surface = pygame.Surface((largeur, hauteur), pygame.SRCALPHA)

    for j in range(rows):
        for i in range(cols):

            if (i < 2 and j < 2) or (i > cols - 3 and j < 2) or \
               (i < 2 and j > rows - 3) or (i > cols - 3 and j > rows - 3):
                continue

            if i == 0 or j == 0 or i == cols - 1 or j == rows - 1:
                color = bord
            elif j < rows // 2.75:
                color = highlight
            else:
                color = couleur

            pygame.draw.rect(
                surface,
                color,
                (i * pixel, j * pixel, pixel, pixel)
            )

    pygame.draw.rect(surface, bord, (1 * pixel, 1 * pixel, pixel, pixel))
    pygame.draw.rect(surface, bord, ((cols - 2) * pixel, 1 * pixel, pixel, pixel))
    pygame.draw.rect(surface, bord, (1 * pixel, (rows - 2) * pixel, pixel, pixel))
    pygame.draw.rect(surface, bord, ((cols - 2) * pixel, (rows - 2) * pixel, pixel, pixel))

    fenetre.blit(surface, (x, y))
    
def dessine_circle(x, y, couleur, fenetre):
    pomme = [
        [0,0,0,3,2,2,0,0],
        [0,0,1,1,1,1,0,0],
        [0,1,1,1,1,1,1,0],
        [0,1,1,1,1,1,1,0],
        [0,1,1,1,1,1,1,0],
        [0,0,1,1,1,1,0,0]
    ]

    largeur = len(pomme[0])
    hauteur = len(pomme)
    taille_totale = int(unite * 1.5)
    pixel_size = taille_totale // largeur

    base_x = x * unite + (unite - taille_totale) // 2
    base_y = y * unite + (unite - taille_totale) // 2

    for j, ligne in enumerate(pomme):
        for i, pixel in enumerate(ligne):
            if pixel == 0:
                continue
            if pixel == 1:
                couleur = couleur_nourriture
            elif pixel == 2:
                couleur = VERT
            elif pixel == 3:
                couleur = MARRON

            pygame.draw.rect(
                fenetre,
                couleur,
                (
                    base_x + i * pixel_size,
                    base_y + j * pixel_size,
                    pixel_size,
                    pixel_size
                )
            )


def dessine_coeur(x, y, couleur, fenetre):
    coeur = [
        [0,1,1,0,0,1,1,0],
        [1,2,2,1,1,2,2,1],
        [1,2,2,2,2,2,2,1],
        [0,1,2,2,2,2,1,0],
        [0,0,1,2,2,1,0,0],
        [0,0,0,1,1,0,0,0]
    ]

    largeur = len(coeur[0])
    hauteur = len(coeur)
    taille_totale = int(unite * 1.25)
    pixel_size = taille_totale // largeur

    base_x = x * unite + (unite - taille_totale) // 2
    base_y = y * unite + (unite - taille_totale) // 2

    for j, ligne in enumerate(coeur):
        for i, pixel in enumerate(ligne):
            if pixel == 0:
                continue
            if pixel == 1:
                couleur = Rouge2
            elif pixel == 2:
                couleur = ROUGE

            pygame.draw.rect(
                fenetre,
                couleur,
                (
                    base_x + i * pixel_size,
                    base_y + j * pixel_size,
                    pixel_size,
                    pixel_size
                )
            )
def snake_body(x, y, fenetre):
    body = [
        [1,1,1,1,1],
        [1,1,1,1,1],
        [1,1,1,1,1],
        [1,1,1,1,1],
        [1,1,1,1,1],
    ]

    taille = unite + 3
    surface = pygame.Surface((taille, taille), pygame.SRCALPHA)

    pixel_size = taille // len(body)

    for j, ligne in enumerate(body):
        for i, pixel in enumerate(ligne):
            if pixel == 0:
                continue
            if pixel == 1:
                couleur = couleur_serpent
            elif pixel == 2:
                couleur = VERT3

            pygame.draw.rect(
                surface,
                couleur,
                (i * pixel_size, j * pixel_size, pixel_size, pixel_size)
            )

    base_x = x * unite + (unite - taille) // 2
    base_y = y * unite + (unite - taille) // 2

    fenetre.blit(surface, (base_x, base_y))

def snake_body_end(x, y, direction, fenetre):
    body_end = [
        [0,2,2,2,0],
        [2,2,2,2,2],
        [2,2,2,2,2],
        [2,2,2,2,2],
        [2,2,2,2,2]
    ]

    taille = unite + 3
    surface_head = pygame.Surface((taille, taille), pygame.SRCALPHA)

    pixel_size = taille // max(len(body_end), len(body_end[0]))

    base_x = x * unite + (unite - taille) // 2
    base_y = y * unite + (unite - taille) // 2

    for j, ligne in enumerate(body_end):
        for i, pixel in enumerate(ligne):
            if pixel == 0:
                continue
            if pixel == 3:
                couleur = VERT3
            elif pixel == 2:
                couleur = couleur_serpent

            pygame.draw.rect(
                surface_head,
                couleur,
                (i * pixel_size, j * pixel_size, pixel_size, pixel_size)
            )

    if direction == "up":
        surface_head = pygame.transform.rotate(surface_head, 180)
    elif direction == "left":
        surface_head = pygame.transform.rotate(surface_head, 270)
    elif direction == "right":
        surface_head = pygame.transform.rotate(surface_head, 90)
    elif direction == "down":
        surface_head = pygame.transform.rotate(surface_head, 0)

    fenetre.blit(surface_head, (base_x, base_y))


def snake_head(x, y, direction, fenetre):
    head = [
        [0,2,2,2,2,2,2,2,2,0],
        [2,2,2,2,2,2,2,2,2,2],
        [2,2,3,3,2,2,3,3,2,2],
        [2,2,3,4,2,2,4,3,2,2],
        [2,2,3,4,2,2,4,3,2,2],
        [0,2,2,2,2,2,2,2,2,0],
        [0,2,2,2,2,2,2,2,2,0],
        [0,0,2,2,5,5,2,2,0,0],
        [0,0,0,0,5,5,0,0,0,0]
    ]

    taille = int(unite * 1.75)
    surface_head = pygame.Surface((taille, taille), pygame.SRCALPHA)

    pixel_size = taille // max(len(head), len(head[0]))

    for j, ligne in enumerate(head):
        for i, pixel in enumerate(ligne):
            if pixel == 0:
                continue
            elif pixel == 2:
                couleur = couleur_serpent
            elif pixel == 3:
                couleur = BLANC
            elif pixel == 4:
                couleur = NOIR
            elif pixel == 5:
                couleur = ROUGE

            pygame.draw.rect(
                surface_head,
                couleur,
                (i * pixel_size, j * pixel_size, pixel_size, pixel_size)
            )

    if direction == "up":
        surface_head = pygame.transform.rotate(surface_head, 180)
        base_x = x * unite + (unite - taille) // 2 - 2.75
        base_y = y * unite + (unite - taille) // 2

    elif direction == "left":
        surface_head = pygame.transform.rotate(surface_head, 270)
        base_x = x * unite + (unite - taille) // 2
        base_y = y * unite + (unite - taille) // 2 + 2.75

    elif direction == "right":
        surface_head = pygame.transform.rotate(surface_head, 90)
        base_x = x * unite + (unite - taille) // 2
        base_y = y * unite + (unite - taille) // 2 - 2.75

    elif direction == "down":
        surface_head = pygame.transform.rotate(surface_head, 0)
        base_x = x * unite + (unite - taille) // 2 + 2.75
        base_y = y * unite + (unite - taille) // 2

    fenetre.blit(surface_head, (base_x, base_y))
    
def snake_turn(x, y, dir_before, dir_after, fenetre):
    turn = [
        [0,0,1,1,1],
        [0,1,1,1,1],
        [1,1,1,1,1],
        [1,1,1,1,1],
        [1,1,1,1,1]
    ]

    taille = unite + 3
    surface = pygame.Surface((taille, taille), pygame.SRCALPHA)

    pixel_size = taille // len(turn)

    for j, ligne in enumerate(turn):
        for i, pixel in enumerate(ligne):
            if pixel == 0:
                continue

            pygame.draw.rect(
                surface,
                couleur_serpent,
                (i * pixel_size, j * pixel_size, pixel_size, pixel_size)
            )

    if dir_before == "up" and dir_after == "right":
        surface = pygame.transform.rotate(surface, 0)

    elif dir_before == "right" and dir_after == "down":
        surface = pygame.transform.rotate(surface, 270)

    elif dir_before == "down" and dir_after == "left":
        surface = pygame.transform.rotate(surface, 180)

    elif dir_before == "left" and dir_after == "up":
        surface = pygame.transform.rotate(surface, 90)

    base_x = x * unite + (unite - taille) // 2
    base_y = y * unite + (unite - taille) // 2

    fenetre.blit(surface, (base_x, base_y))
    
def snake_turn_otherside(x, y, dir_before, dir_after, fenetre):
    turn_otherside = [
        [1,1,1,0,0],
        [1,1,1,1,0],
        [1,1,1,1,1],
        [1,1,1,1,1],
        [1,1,1,1,1]
    ]

    taille = unite + 3
    surface = pygame.Surface((taille, taille), pygame.SRCALPHA)

    pixel_size = taille // len(turn_otherside)

    for j, ligne in enumerate(turn_otherside):
        for i, pixel in enumerate(ligne):
            if pixel == 0:
                continue

            pygame.draw.rect(
                surface,
                couleur_serpent,
                (i * pixel_size, j * pixel_size, pixel_size, pixel_size)
            )


    if dir_before == "up" and dir_after == "left":
        surface = pygame.transform.rotate(surface, 0)

    elif dir_before == "left" and dir_after == "down":
        surface = pygame.transform.rotate(surface, 90)

    elif dir_before == "down" and dir_after == "right":
        surface = pygame.transform.rotate(surface, 180)

    elif dir_before == "right" and dir_after == "up":
        surface = pygame.transform.rotate(surface, 270)

    base_x = x * unite + (unite - taille) // 2
    base_y = y * unite + (unite - taille) // 2

    fenetre.blit(surface, (base_x, base_y))
    
def dessine_coeur_score(x, y, couleur, fenetre):
    coeur_score = [
        [0,1,1,0,0,1,1,0],
        [1,1,1,1,1,1,1,1],
        [1,1,1,1,1,1,1,1],
        [0,1,1,1,1,1,1,0],
        [0,0,1,1,1,1,0,0],
        [0,0,0,1,1,0,0,0]
    ]
    largeur = len(coeur_score[0])
    hauteur = len(coeur_score)
    taille_totale = int(unite * 3)
    pixel_size = taille_totale // largeur
    base_x = x * unite + (unite - taille_totale) // 2
    base_y = y * unite + (unite - taille_totale) // 2

    for j, ligne in enumerate(coeur_score):
        for i, pixel in enumerate(ligne):
            if pixel == 1:
                pygame.draw.rect(
                    fenetre,
                    couleur,
                    (base_x + i * pixel_size, base_y + j * pixel_size, pixel_size, pixel_size)
                )
            
def rotation_aleatoire(pos):
    return [0, 90, 180, 270][(pos[0] * 31 + pos[1] * 17) % 4]
            
def dessine_plusieurs_carres(positions, direction, fenetre):
    for i in range(1, len(positions) - 1):
        x, y = positions[i]

        dir_before = get_dir(positions[i - 1], positions[i])
        dir_after = get_dir(positions[i], positions[i + 1])

        if dir_before != dir_after:
            
            if (dir_before == "up" and dir_after == "right") or \
                (dir_before == "right" and dir_after == "down") or \
                (dir_before == "down" and dir_after == "left") or \
                (dir_before == "left" and dir_after == "up"):

                snake_turn(x, y, dir_before, dir_after, fenetre)
                
            else:
                snake_turn_otherside(x, y, dir_before, dir_after, fenetre)
                
        else:
            snake_body(x, y, fenetre)

    x, y = positions[-1]
    snake_head(x, y, direction, fenetre)

    snake_tail(positions, fenetre)
    
def snake_tail(positions, fenetre):
    tail = positions[0]
    after_tail = positions[1]

    direction = get_dir(tail, after_tail)
    snake_body_end(tail[0], tail[1], direction, fenetre)


def entier_aleatoire(min_v, max_v):
    return random.randint(min_v, max_v)


def position_aleatoire():
    return (
        random.randint(13, largeur_zone - 14),
        random.randint(1, hauteur_zone - 2)
    )


def afficher_message_score(message: str, couleur: tuple, fenetre):
    BASE_DIR = os.path.dirname(__file__)
    chemin_font = os.path.join(BASE_DIR, "assets", "font.otf")
    myfont = pygame.font.Font(chemin_font, 100)
    texte = myfont.render(message, True, couleur)
    fenetre.blit(texte, (25, 0))


def afficher_message_record(message: str, couleur: tuple, fenetre):
    BASE_DIR = os.path.dirname(__file__)
    chemin_font = os.path.join(BASE_DIR, "assets", "font.otf")
    myfont = pygame.font.Font(chemin_font, 100)
    texte = myfont.render(message, True, couleur)
    fenetre.blit(texte, (1510, 0))


def afficher_message_gameover(message: str, couleur: tuple, fenetre):
    BASE_DIR = os.path.dirname(__file__)
    chemin_font = os.path.join(BASE_DIR, "assets", "font.otf")
    myfont = pygame.font.Font(chemin_font, 200)
    texte = myfont.render(message, True, couleur)
    fenetre.blit(texte, (575, 450))


def afficher_message_coeur(message: str, couleur: tuple, fenetre):
    BASE_DIR = os.path.dirname(__file__)
    chemin_font = os.path.join(BASE_DIR, "assets", "font.otf")
    myfont = pygame.font.Font(chemin_font, 100)
    texte = myfont.render(message, True, couleur)
    fenetre.blit(texte, (25, 875))
    
def afficher_nom(message, couleur, fenetre):
    BASE_DIR = os.path.dirname(__file__)
    chemin_font = os.path.join(BASE_DIR, "assets", "font.otf")
    myfont = pygame.font.Font(chemin_font, 70)
    texte = myfont.render(message, True, couleur)
    fenetre.blit(texte, (25, 10))
    
def afficher_nom_ingame(message, couleur, fenetre):
    BASE_DIR = os.path.dirname(__file__)
    chemin_font = os.path.join(BASE_DIR, "assets", "font.otf")
    myfont = pygame.font.Font(chemin_font, 70)
    texte = myfont.render(message, True, couleur)
    fenetre.blit(texte, (25, 75))
    
def afficher_message_best(message: str, couleur: tuple, fenetre):
    BASE_DIR = os.path.dirname(__file__)
    chemin_font = os.path.join(BASE_DIR, "assets", "font.otf")
    myfont = pygame.font.Font(chemin_font, 100)
    texte = myfont.render(message, True, couleur)
    fenetre.blit(texte, (1510, 75))


def nouvelle_direction(touche, direction):
    match touche:
        case pygame.K_z:
            if direction != "down":
                return "up"
        case pygame.K_s:
            if direction != "up":
                return "down"
        case pygame.K_q:
            if direction != "right":
                return "left"
        case pygame.K_d:
            if direction != "left":
                return "right"
        case pygame.K_UP:
            if direction != "down":
                return "up"
        case pygame.K_DOWN:
            if direction != "up":
                return "down"
        case pygame.K_LEFT:
            if direction != "right":
                return "left"
        case pygame.K_RIGHT:
            if direction != "left":
                return "right"
    return direction


def rebond_direction(pos, direction):
    x, y = pos

    if x < LIM_LEFT and direction == "left":
        return "up"
    if x > LIM_RIGHT and direction == "right":
        return "down"
    if y < LIM_TOP and direction == "up":
        return "right"
    if y > LIM_BOTTOM and direction == "down":
        return "left"

    return direction


def get_dir(a, b):
    if a[0] < b[0]:
        return "right"
    if a[0] > b[0]:
        return "left"
    if a[1] < b[1]:
        return "down"
    if a[1] > b[1]:
        return "up"
    
def direction_inverse(dir1, dir2):
    return (dir1 == "up" and dir2 == "down") or \
        (dir1 == "down" and dir2 == "up") or \
        (dir1 == "left" and dir2 == "right") or \
        (dir1 == "right" and dir2 == "left")


def nouvelle_position(pos: tuple, direction: str):
    match direction:
        case "up":
            return (pos[0], pos[1] - 1)
        case "down":
            return (pos[0], pos[1] + 1)
        case "right":
            return (pos[0] + 1, pos[1])
        case "left":
            return (pos[0] - 1, pos[1])
    return pos

def is_fed(pos, food_pos):
    return pos[-1] == food_pos


def is_fed_coeur(pos, position_coeur):
    return pos[-1] == position_coeur


def is_touching(positions_serpent: list, position_tete):
    for pos in positions_serpent:
        if pos == position_tete:
            return True
    return False


def is_on_board(pos):
    x, y = pos
    return LIM_LEFT <= x <= LIM_RIGHT and LIM_TOP <= y <= LIM_BOTTOM


def tour_de_jeu(
    positions_serpent: list,
    jeu_fini: bool,
    direction_serpent: tuple,
    position_pomme: tuple,
    position_coeur: tuple,
    score_coeur,
    prochain_coeur,
    fenetre,
    nom_joueur,
    record_global,
) -> tuple[list, bool, tuple, tuple, tuple, int, int, int, tuple, int]:
    dessine_circle(position_pomme[0], position_pomme[1], couleur_nourriture, fenetre)

    score = len(positions_serpent) - 5
    afficher_message_score("Score : " + str(score), BLANC, fenetre)
    afficher_message_record("Record : " + str(meilleur_score), BLEU2, fenetre)
    afficher_message_coeur("HP : " + str(score_coeur), NOIR, fenetre)
    afficher_nom_ingame(nom_joueur, NOIR, fenetre)
    afficher_message_best("Best      : " + str(record_global), VERT4, fenetre)

    dessine_plusieurs_carres(
        positions_serpent,
        direction_serpent,
        fenetre
    )   

    if position_coeur is not None:
        dessine_coeur(position_coeur[0], position_coeur[1], couleur_coeur, fenetre)

    pygame.display.update()

    nouvelle = direction_serpent

    for event in pygame.event.get():
        if event.type == pygame.QUIT:
            jeu_fini = True

        elif event.type == pygame.KEYDOWN:
            tentative = nouvelle_direction(event.key, direction_serpent)
            if not direction_inverse(direction_serpent, tentative):
                nouvelle = tentative


    direction_serpent = nouvelle

    head = positions_serpent[-1]
    future_pos = nouvelle_position(head, direction_serpent)
    
    if not is_on_board(future_pos):
        score_coeur -= 1
        if score_coeur == 0:
            jeu_fini = True
        else:
            direction_serpent = rebond_direction(future_pos, direction_serpent)
            future_pos = nouvelle_position(head, direction_serpent)

    if future_pos != position_pomme:
        positions_serpent.pop(0)
    else:
        position_pomme = position_aleatoire()

    positions_serpent.append(future_pos)


    if position_coeur is not None and future_pos == position_coeur:
        position_coeur = None
        if score_coeur < 3:
            score_coeur += 1

    if is_touching(positions_serpent[:-1], future_pos):
        score_coeur -= 1
        if score_coeur == 0:
            jeu_fini = True

    return (
        positions_serpent,
        jeu_fini,
        direction_serpent,
        position_pomme,
        position_coeur,
        score_coeur,
        score,
        prochain_coeur,
    )

def menu_accueil(fenetre, nom_joueur, meilleur_score, record_global):
    BASE_DIR = os.path.dirname(__file__)
    chemin_font = os.path.join(BASE_DIR, "assets", "font.otf")

    titre_font = pygame.font.Font(chemin_font, 180)
    bouton_font = pygame.font.Font(chemin_font, 100)

    bouton_largeur = 400
    bouton_hauteur = 150

    bouton_x = (largeur_zone * unite - bouton_largeur) // 2
    bouton_y = (hauteur_zone * unite - bouton_hauteur) // 2 + 100

    bouton_rect = pygame.Rect(
        bouton_x,
        bouton_y,
        bouton_largeur,
        bouton_hauteur
    )

    dans_menu = True

    while dans_menu:
        fenetre.fill(couleur_fond)
        afficher_message_record("Record : " + str(meilleur_score), BLEU2, fenetre)
        afficher_message_best("Best      : " + str(record_global), VERT4, fenetre)
        afficher_nom(nom_joueur, NOIR, fenetre)

        titre = titre_font.render("SNAKE.PY", True, VERT3)
        fenetre.blit(
            titre,
            (
                (largeur_zone * unite - titre.get_width()) // 2,
                180
            )
        )

        souris = pygame.mouse.get_pos()

        if bouton_rect.collidepoint(souris):
            couleur_bouton = VERT3
        else:
            couleur_bouton = VERT

        dessine_bouton_pixel(
            bouton_x,
            bouton_y,
            bouton_largeur,
            bouton_hauteur,
            couleur_bouton,
            fenetre
        )

        texte_bouton = bouton_font.render("START", True, BLANC)
        fenetre.blit(
            texte_bouton,
            (
                bouton_x + (bouton_largeur - texte_bouton.get_width()) // 2,
                bouton_y + (bouton_hauteur - texte_bouton.get_height()) // 2
            )
        )

        pygame.display.update()

        for event in pygame.event.get():
            if event.type == pygame.QUIT:
                pygame.quit()
                quit()

            if event.type == pygame.MOUSEBUTTONDOWN:
                if bouton_rect.collidepoint(event.pos):
                    dans_menu = False

            if event.type == pygame.KEYDOWN:
                if event.key == pygame.K_RETURN:
                    dans_menu = False

                if event.key == pygame.K_ESCAPE:
                    nom_joueur = demander_nom_joueur(fenetre)

                    meilleur_score = charger_record_joueur(nom_joueur)
                    record_global = charger_record_global()
                    
    return nom_joueur
            
def demander_nom_joueur(fenetre):
    BASE_DIR = os.path.dirname(__file__)
    chemin_font = os.path.join(BASE_DIR, "assets", "font.otf")

    titre_font = pygame.font.Font(chemin_font, 100)
    texte_font = pygame.font.Font(chemin_font, 70)

    nom = ""
    saisie = True

    while saisie:
        fenetre.fill(couleur_fond)

        titre = titre_font.render("Enter your username", True, VERT3)
        fenetre.blit(
            titre,
            (
                (largeur_zone * unite - titre.get_width()) // 2,
                250
            )
        )

        pygame.draw.rect(
            fenetre,
            VERT,
            (500, 500, 900, 100),
            4
        )

        texte_nom = texte_font.render(nom, True, NOIR)
        fenetre.blit(texte_nom, (530, 515))

        pygame.display.update()

        for event in pygame.event.get():
            if event.type == pygame.QUIT:
                pygame.quit()
                quit()

            if event.type == pygame.KEYDOWN:

                if event.key == pygame.K_RETURN and nom != "":
                    saisie = False

                elif event.key == pygame.K_BACKSPACE:
                    nom = nom[:-1]

                elif event.key == pygame.K_ESCAPE:
                    pygame.quit()
                    quit()

                else:
                    if len(nom) < 12:
                        nom += event.unicode

    return nom

def ecran_game_over(fenetre, score, meilleur_score, nom_joueur):
    attente = True

    while attente:
        fenetre.fill(NOIR)

        afficher_message_gameover("GAME OVER", ROUGE, fenetre)
        afficher_message_score("Score : " + str(score), BLANC, fenetre)
        afficher_message_record("Record : " + str(meilleur_score), BLEU2, fenetre)

        pygame.display.update()

        for event in pygame.event.get():
            if event.type == pygame.QUIT:
                attente = False

            if event.type == pygame.KEYDOWN:
                if event.key == pygame.K_RETURN:
                    start(nom_joueur, record_global)

                if event.key == pygame.K_ESCAPE:
                    attente = False


def start(nom_joueur, record_global):
    global meilleur_score
    meilleur_score = charger_record_joueur(nom_joueur)
    
    positions_serpent = [(26, 10), (26, 11), (25, 11), (24, 11), (23, 11)]
    direction_serpent = "left"
    position_pomme = position_aleatoire()
    
    rotations_body = []

    position_coeur = None
    prochain_coeur = 10

    pygame.init()
    clock = pygame.time.Clock()

    fenetre = pygame.display.set_mode(
        (largeur_zone * unite, hauteur_zone * unite),
        pygame.FULLSCREEN
    )

    pygame.display.set_caption("Snake")

    nom_joueur = menu_accueil(
        fenetre,
        nom_joueur,
        meilleur_score,
        record_global
    )

    jeu_fini = False
    
    score_coeur = 1

    tick = 10

    while jeu_fini == False:
        fenetre.fill(couleur_fond)

        dessine_rect(0, 0, VERT3, fenetre)
        dessine_rect(13, 0, couleur_fond, fenetre)
        dessine_rect(47, 0, VERT3, fenetre)

        score = len(positions_serpent) - 5

        if score >= prochain_coeur and position_coeur is None:
            position_coeur = position_aleatoire()
            while ( position_coeur in positions_serpent or not is_on_board(position_coeur) ):
                position_coeur = position_aleatoire()
            prochain_coeur += 10
                
        if position_coeur is not None and is_fed_coeur(positions_serpent, position_coeur):
            position_coeur = None
            
        if score_coeur == 1:
            dessine_coeur_score(1.85, 32, couleur_coeur, fenetre)
            dessine_coeur_score(5.85, 32, NOIR, fenetre)
            dessine_coeur_score(9.85, 32, NOIR, fenetre)
            
        elif score_coeur == 2:
            dessine_coeur_score(1.85, 32, couleur_coeur, fenetre)
            dessine_coeur_score(5.85, 32, couleur_coeur, fenetre)
            dessine_coeur_score(9.85, 32, NOIR, fenetre)
            
        elif score_coeur == 3:
            dessine_coeur_score(1.85, 32, couleur_coeur, fenetre)
            dessine_coeur_score(5.85, 32, couleur_coeur, fenetre)
            dessine_coeur_score(9.85, 32, couleur_coeur, fenetre)

        (
            positions_serpent,
            jeu_fini,
            direction_serpent,
            position_pomme,
            position_coeur,
            score_coeur,
            score,
            prochain_coeur,
        ) = tour_de_jeu(
            positions_serpent,
            jeu_fini,
            direction_serpent,
            position_pomme,
            position_coeur,
            score_coeur,
            prochain_coeur,
            fenetre,
            nom_joueur,
            record_global,
        )
        
        clock.tick(tick)
        tick = len(positions_serpent) / 2 + 2

        if score > meilleur_score:
            meilleur_score = score
            sauvegarder_record_joueur(nom_joueur, meilleur_score)
            
            record_global = charger_record_global()

    ecran_game_over(fenetre, score, meilleur_score, nom_joueur)

    pygame.quit()
    quit()


pygame.init()

fenetre = pygame.display.set_mode(
    (largeur_zone * unite, hauteur_zone * unite),
    pygame.FULLSCREEN
)

nom_joueur = demander_nom_joueur(fenetre)

meilleur_score = charger_record_joueur(nom_joueur)
record_global = charger_record_global()

nom_joueur = menu_accueil(
    fenetre,
    nom_joueur,
    meilleur_score,
    record_global
)

start(nom_joueur, record_global)